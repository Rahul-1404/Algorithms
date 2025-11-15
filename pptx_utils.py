"""
PowerPoint Utilities
Validation, thumbnail generation, and helper functions
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image, ImageDraw
from io import BytesIO
from typing import List, Dict, Optional, Tuple
import zipfile
import os
from pathlib import Path
import xml.etree.ElementTree as ET


class Validator:
    """Validate PowerPoint presentations"""

    @staticmethod
    def validate_presentation(pptx_path: str) -> Dict:
        """
        Validate PowerPoint file

        Returns:
            Dict with validation results
        """
        results = {
            'valid': True,
            'errors': [],
            'warnings': [],
            'info': {}
        }

        try:
            prs = Presentation(pptx_path)

            # Basic info
            results['info']['slide_count'] = len(prs.slides)
            results['info']['slide_width'] = prs.slide_width.inches
            results['info']['slide_height'] = prs.slide_height.inches

            # Check each slide
            for idx, slide in enumerate(prs.slides):
                # Check for text overflow
                for shape in slide.shapes:
                    if hasattr(shape, 'text_frame'):
                        if shape.text_frame.text and len(shape.text_frame.text) > 1000:
                            results['warnings'].append(
                                f"Slide {idx + 1}: Shape '{shape.name}' may have too much text"
                            )

                        # Check for very small text
                        for paragraph in shape.text_frame.paragraphs:
                            for run in paragraph.runs:
                                if run.font.size and run.font.size.pt < 10:
                                    results['warnings'].append(
                                        f"Slide {idx + 1}: Very small font size ({run.font.size.pt}pt) detected"
                                    )

                # Check for missing images
                for shape in slide.shapes:
                    if shape.shape_type == 13:  # Picture
                        try:
                            img = shape.image
                        except Exception as e:
                            results['errors'].append(
                                f"Slide {idx + 1}: Missing or corrupted image"
                            )
                            results['valid'] = False

        except Exception as e:
            results['valid'] = False
            results['errors'].append(f"Failed to open presentation: {str(e)}")

        return results

    @staticmethod
    def check_file_size(pptx_path: str, max_size_mb: float = 50) -> Dict:
        """
        Check file size and embedded resources

        Args:
            pptx_path: Path to .pptx file
            max_size_mb: Maximum recommended size in MB
        """
        results = {
            'file_size_mb': 0,
            'over_limit': False,
            'embedded_images': [],
            'warnings': []
        }

        file_size = os.path.getsize(pptx_path) / (1024 * 1024)
        results['file_size_mb'] = round(file_size, 2)
        results['over_limit'] = file_size > max_size_mb

        # Check embedded resources
        try:
            with zipfile.ZipFile(pptx_path, 'r') as zip_ref:
                for file in zip_ref.namelist():
                    if file.startswith('ppt/media/'):
                        file_info = zip_ref.getinfo(file)
                        size_kb = file_info.file_size / 1024

                        results['embedded_images'].append({
                            'filename': file,
                            'size_kb': round(size_kb, 2)
                        })

                        if size_kb > 1024:  # > 1MB
                            results['warnings'].append(
                                f"Large embedded file: {file} ({round(size_kb/1024, 2)}MB)"
                            )

        except Exception as e:
            results['warnings'].append(f"Could not analyze embedded files: {str(e)}")

        return results


class ThumbnailGenerator:
    """Generate thumbnail previews of slides"""

    @staticmethod
    def generate_slide_thumbnails(pptx_path: str, output_dir: str,
                                  width: int = 320, height: int = 180) -> List[str]:
        """
        Generate thumbnail images for all slides

        Args:
            pptx_path: Path to .pptx file
            output_dir: Directory to save thumbnails
            width, height: Thumbnail dimensions

        Returns:
            List of generated thumbnail paths
        """
        # Note: python-pptx doesn't directly support rendering to images
        # This is a simplified version that creates placeholder thumbnails
        # For real rendering, you'd need LibreOffice or similar

        Path(output_dir).mkdir(parents=True, exist_ok=True)

        prs = Presentation(pptx_path)
        thumbnail_paths = []

        for idx, slide in enumerate(prs.slides):
            # Create simplified thumbnail
            thumbnail = ThumbnailGenerator._create_slide_thumbnail(
                slide, prs.slide_width.inches, prs.slide_height.inches,
                width, height
            )

            output_path = os.path.join(output_dir, f'slide_{idx + 1}.png')
            thumbnail.save(output_path)
            thumbnail_paths.append(output_path)

        return thumbnail_paths

    @staticmethod
    def _create_slide_thumbnail(slide, slide_width: float, slide_height: float,
                               thumb_width: int, thumb_height: int) -> Image.Image:
        """Create simplified thumbnail of slide"""
        # Create blank thumbnail
        thumbnail = Image.new('RGB', (thumb_width, thumb_height), (255, 255, 255))
        draw = ImageDraw.Draw(thumbnail)

        # Add border
        draw.rectangle([0, 0, thumb_width - 1, thumb_height - 1],
                      outline=(200, 200, 200), width=2)

        # Try to extract basic info
        shape_count = len(slide.shapes)
        text_shapes = sum(1 for shape in slide.shapes if hasattr(shape, 'text'))

        # Draw simple representation
        draw.text((10, 10), f"Shapes: {shape_count}", fill=(0, 0, 0))
        draw.text((10, 30), f"Text: {text_shapes}", fill=(0, 0, 0))

        # Try to represent layout
        scale_x = thumb_width / (slide_width * 72)
        scale_y = thumb_height / (slide_height * 72)

        for shape in slide.shapes:
            try:
                left = int(shape.left.pt * scale_x)
                top = int(shape.top.pt * scale_y)
                width = int(shape.width.pt * scale_x)
                height = int(shape.height.pt * scale_y)

                # Draw shape outline
                if hasattr(shape, 'text') and shape.text:
                    color = (100, 100, 200)  # Blue for text
                else:
                    color = (200, 100, 100)  # Red for other shapes

                draw.rectangle([left, top, left + width, top + height],
                              outline=color, width=1)
            except:
                pass

        return thumbnail

    @staticmethod
    def create_slide_grid(pptx_path: str, output_path: str,
                         columns: int = 3) -> str:
        """
        Create a grid view of all slides

        Args:
            pptx_path: Path to .pptx file
            output_path: Output image path
            columns: Number of columns in grid

        Returns:
            Path to generated grid image
        """
        prs = Presentation(pptx_path)
        slide_count = len(prs.slides)

        thumb_width = 320
        thumb_height = 180
        padding = 20

        rows = (slide_count + columns - 1) // columns

        grid_width = columns * thumb_width + (columns + 1) * padding
        grid_height = rows * thumb_height + (rows + 1) * padding

        grid = Image.new('RGB', (grid_width, grid_height), (240, 240, 240))

        for idx, slide in enumerate(prs.slides):
            row = idx // columns
            col = idx % columns

            x = col * thumb_width + (col + 1) * padding
            y = row * thumb_height + (row + 1) * padding

            thumbnail = ThumbnailGenerator._create_slide_thumbnail(
                slide, prs.slide_width.inches, prs.slide_height.inches,
                thumb_width, thumb_height
            )

            grid.paste(thumbnail, (x, y))

            # Add slide number
            draw = ImageDraw.Draw(grid)
            draw.text((x + 5, y + 5), f"#{idx + 1}", fill=(255, 0, 0))

        grid.save(output_path)
        return output_path


class PPTXHelper:
    """Helper utilities for PowerPoint operations"""

    @staticmethod
    def merge_presentations(pptx_files: List[str], output_path: str):
        """
        Merge multiple PowerPoint files into one

        Args:
            pptx_files: List of .pptx file paths
            output_path: Output file path
        """
        if not pptx_files:
            raise ValueError("No files to merge")

        # Start with first presentation
        merged = Presentation(pptx_files[0])

        # Add slides from other presentations
        for pptx_file in pptx_files[1:]:
            prs = Presentation(pptx_file)

            for slide in prs.slides:
                # This is complex - simplified version
                # In practice, you'd need to copy all shapes and properties
                merged.slides.add_slide(slide.slide_layout)

        merged.save(output_path)

    @staticmethod
    def extract_notes(pptx_path: str) -> Dict[int, str]:
        """
        Extract speaker notes from slides

        Args:
            pptx_path: Path to .pptx file

        Returns:
            Dict mapping slide index to notes text
        """
        prs = Presentation(pptx_path)
        notes = {}

        for idx, slide in enumerate(prs.slides):
            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
                text_frame = notes_slide.notes_text_frame
                notes[idx] = text_frame.text

        return notes

    @staticmethod
    def add_slide_numbers(pptx_path: str, output_path: str, start_number: int = 1):
        """
        Add slide numbers to presentation

        Args:
            pptx_path: Input .pptx path
            output_path: Output .pptx path
            start_number: Starting slide number
        """
        prs = Presentation(pptx_path)

        for idx, slide in enumerate(prs.slides):
            # Add slide number in bottom right
            slide_num = start_number + idx

            textbox = slide.shapes.add_textbox(
                prs.slide_width - Inches(1),
                prs.slide_height - Inches(0.5),
                Inches(0.5),
                Inches(0.3)
            )

            text_frame = textbox.text_frame
            text_frame.text = str(slide_num)
            paragraph = text_frame.paragraphs[0]
            paragraph.font.size = Pt(12)

            textbox.line.fill.background()

        prs.save(output_path)

    @staticmethod
    def extract_text(pptx_path: str) -> List[Dict]:
        """
        Extract all text from presentation

        Args:
            pptx_path: Path to .pptx file

        Returns:
            List of dicts with slide index and text
        """
        prs = Presentation(pptx_path)
        text_data = []

        for idx, slide in enumerate(prs.slides):
            slide_text = []

            for shape in slide.shapes:
                if hasattr(shape, 'text'):
                    slide_text.append(shape.text)

            text_data.append({
                'slide': idx + 1,
                'text': '\n'.join(slide_text)
            })

        return text_data

    @staticmethod
    def get_presentation_stats(pptx_path: str) -> Dict:
        """
        Get comprehensive presentation statistics

        Args:
            pptx_path: Path to .pptx file

        Returns:
            Dict with statistics
        """
        prs = Presentation(pptx_path)

        stats = {
            'slide_count': len(prs.slides),
            'slide_width_inches': prs.slide_width.inches,
            'slide_height_inches': prs.slide_height.inches,
            'aspect_ratio': f"{round(prs.slide_width.inches / prs.slide_height.inches, 2)}:1",
            'total_shapes': 0,
            'total_text_boxes': 0,
            'total_images': 0,
            'total_charts': 0,
            'total_tables': 0,
        }

        for slide in prs.slides:
            stats['total_shapes'] += len(slide.shapes)

            for shape in slide.shapes:
                if hasattr(shape, 'text'):
                    stats['total_text_boxes'] += 1
                if shape.shape_type == 13:  # Picture
                    stats['total_images'] += 1
                if shape.shape_type == 3:  # Chart
                    stats['total_charts'] += 1
                if shape.shape_type == 19:  # Table
                    stats['total_tables'] += 1

        return stats


if __name__ == '__main__':
    # Demo utilities
    print("PowerPoint Utilities Demo")

    # Create a sample presentation for testing
    from pptx_templates import ModernTemplates

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)

    ModernTemplates.create_title_slide(prs, "Test Presentation", "For Utils Demo", (79, 70, 229))
    ModernTemplates.create_content_slide(prs, "Sample Slide", ["Point 1", "Point 2", "Point 3"], (79, 70, 229))

    test_file = 'utils_test.pptx'
    prs.save(test_file)

    # Validate
    validation = Validator.validate_presentation(test_file)
    print(f"\nValidation: {validation}")

    # Get stats
    stats = PPTXHelper.get_presentation_stats(test_file)
    print(f"\nStats: {stats}")

    # Generate thumbnails
    thumbnails = ThumbnailGenerator.generate_slide_thumbnails(test_file, 'thumbnails')
    print(f"\nThumbnails generated: {len(thumbnails)}")

    # Create grid
    grid_path = ThumbnailGenerator.create_slide_grid(test_file, 'slide_grid.png')
    print(f"Grid created: {grid_path}")

    print("\nUtils demo complete!")
