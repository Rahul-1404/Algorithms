#!/usr/bin/env python3
"""
Quick Start Script for Python PowerPoint Creation
Demonstrates the easiest way to create beautiful presentations
"""

from pptx import Presentation
from pptx.util import Inches
from html2pptx import create_presentation_from_html
from pptx_charts import ChartGenerator
from pptx_templates import ModernTemplates


def create_quick_presentation():
    """Create a simple but beautiful presentation in minutes"""

    print("Creating quick start presentation...")

    # Method 1: Using HTML (easiest for web developers)
    print("\n1. Creating slides from HTML...")

    html_slides = [
        # Slide 1: Title with gradient
        """
        <div style="width: 720pt; height: 405pt; background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);">
            <h1 style="position: absolute; left: 50pt; top: 150pt; width: 620pt; color: white;
                       font-size: 54pt; font-family: Arial; text-align: center; font-weight: bold;">
                Quick Start Demo
            </h1>
            <p style="position: absolute; left: 50pt; top: 250pt; width: 620pt; color: white;
                      font-size: 28pt; font-family: Arial; text-align: center;">
                Beautiful PowerPoint with Python
            </p>
        </div>
        """,

        # Slide 2: Content slide
        """
        <div style="width: 720pt; height: 405pt; background: white;">
            <div style="position: absolute; left: 0; top: 0; width: 720pt; height: 80pt;
                       background: #4f46e5;"></div>

            <h2 style="position: absolute; left: 30pt; top: 20pt; color: white;
                       font-size: 36pt; font-family: Arial; font-weight: bold;">
                Why Use Python for PowerPoint?
            </h2>

            <p style="position: absolute; left: 50pt; top: 120pt; width: 620pt;
                      font-size: 20pt; font-family: Arial; color: #374151;">
                ✓ Pure Python - no Node.js required<br/>
                ✓ Beautiful gradients and modern designs<br/>
                ✓ Professional charts with matplotlib<br/>
                ✓ Easy template system<br/>
                ✓ Works anywhere Python runs
            </p>
        </div>
        """
    ]

    create_presentation_from_html(html_slides, 'quickstart_html.pptx')
    print("   ✓ Created: quickstart_html.pptx")

    # Method 2: Using templates (easiest for non-coders)
    print("\n2. Creating slides from templates...")

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)

    color = (79, 70, 229)  # Modern indigo

    # Add title slide
    ModernTemplates.create_title_slide(
        prs, "Template Demo", "Professional and Easy", color
    )

    # Add content slide
    ModernTemplates.create_content_slide(
        prs, "Key Features",
        [
            "Modern, professional designs",
            "Multiple layout options",
            "Consistent styling",
            "Easy to customize"
        ],
        color, 'two_column'
    )

    # Add quote slide
    ModernTemplates.create_quote_slide(
        prs,
        "Simplicity is the ultimate sophistication",
        "Leonardo da Vinci",
        color
    )

    prs.save('quickstart_templates.pptx')
    print("   ✓ Created: quickstart_templates.pptx")

    # Method 3: Adding charts (for data presentations)
    print("\n3. Creating presentation with charts...")

    prs2 = Presentation()
    prs2.slide_width = Inches(10)
    prs2.slide_height = Inches(5.625)

    # Title
    ModernTemplates.create_title_slide(
        prs2, "Data Visualization", "Charts Made Easy", (236, 72, 153)
    )

    # Chart slide
    blank_layout = prs2.slide_layouts[6]
    slide = prs2.slides.add_slide(blank_layout)

    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    # Add title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_box.text_frame
    title_frame.text = "Sales Performance"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(32)
    title_para.font.bold = True
    title_box.line.fill.background()

    # Add chart
    ChartGenerator.add_matplotlib_chart(
        slide, 1, 1.2, 8, 4,
        ChartGenerator.create_beautiful_bar_chart,
        categories=['Jan', 'Feb', 'Mar', 'Apr'],
        data_series=[
            {'name': '2023', 'values': [100, 120, 110, 130]},
            {'name': '2024', 'values': [130, 150, 145, 170]}
        ],
        title='',
        color_scheme='modern'
    )

    prs2.save('quickstart_charts.pptx')
    print("   ✓ Created: quickstart_charts.pptx")

    print("\n" + "=" * 60)
    print("Quick start complete! Created 3 presentations:")
    print("  1. quickstart_html.pptx - HTML-based slides")
    print("  2. quickstart_templates.pptx - Template-based slides")
    print("  3. quickstart_charts.pptx - Data visualization")
    print("=" * 60)
    print("\nNext steps:")
    print("  • Run 'python examples.py' for comprehensive examples")
    print("  • Check README.md for detailed documentation")
    print("  • Explore individual modules for advanced features")


if __name__ == '__main__':
    create_quick_presentation()
